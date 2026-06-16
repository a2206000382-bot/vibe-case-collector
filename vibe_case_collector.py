#!/usr/bin/env python3
"""
Vibe Coding independent startup case collector.

The script searches public web pages, extracts verifiable case information with
an OpenAI-compatible chat API, enforces a daily CNY token budget, and writes TXT
and Word archives with the fixed Chinese template requested by the operator.
"""

from __future__ import annotations

import argparse
import dataclasses
import datetime as dt
import json
import os
import re
import sys
import time
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional
from urllib.parse import urlparse

import requests
from bs4 import BeautifulSoup
from docx import Document
from dotenv import load_dotenv

try:
    from ddgs import DDGS
except ImportError:  # pragma: no cover - compatibility for older installs
    from duckduckgo_search import DDGS  # type: ignore


MISSING = "未披露"
INFERENCE_MARKERS = ("推断", "猜测", "可能", "似乎", "应该", "疑似")
TRUST_LEVELS = {"★", "★★", "★★★", "★★★★"}
CODING_BACKGROUNDS = {"有", "无", "自学", MISSING}
PRODUCT_TYPES = {
    "SaaS订阅工具",
    "本地客户端",
    "网页插件",
    "API服务",
    "开源免费工具",
    MISSING,
}
INCOME_MODE_KEYWORDS = (
    "SaaS月订阅",
    "一次性买断",
    "API按量收费",
    "广告变现",
    "企业定制服务",
)
RISK_KEYS = ("合规风险", "平台依赖风险", "市场竞争风险", "长期运营风险")
OPPORTUNITY_KEYS = ("横向扩展场景", "纵向功能深化", "B端企业转化", "细分生态位卡位")
CASE_FIELDS = (
    "case_id",
    "product_name",
    "founder_background",
    "coding_background",
    "product_type",
    "product_usage",
    "target_users",
    "pain_points",
    "income_model",
    "monthly_income",
    "ai_tools_used",
    "development_time",
    "data_sources",
    "credibility",
    "solo_project",
    "risks",
    "opportunities",
    "evidence_summary",
    "missing_fields",
    "completeness_score",
    "verification_notes",
)


@dataclasses.dataclass
class Config:
    keywords: List[str]
    daily_budget_cny: float
    llm_api_key: str
    llm_base_url: str
    llm_model: str
    llm_input_price_cny_per_1m_tokens: float
    llm_output_price_cny_per_1m_tokens: float
    search_provider: str
    serper_api_key: str
    max_results_per_keyword: int
    max_pages_per_keyword: int
    max_source_chars: int
    max_llm_output_tokens: int
    request_timeout_seconds: int
    output_dir: Path
    run_time: str
    exchange_rates: str
    enable_enrichment: bool
    enrichment_results_per_case: int
    enrichment_source_chars: int
    min_completeness_score: float
    rejected_output: bool


def getenv_str(name: str, default: str = "") -> str:
    return os.getenv(name, default).strip()


def getenv_int(name: str, default: int) -> int:
    raw = getenv_str(name, str(default))
    try:
        return int(raw)
    except ValueError:
        return default


def getenv_float(name: str, default: float) -> float:
    raw = getenv_str(name, str(default))
    try:
        return float(raw)
    except ValueError:
        return default


def getenv_bool(name: str, default: bool) -> bool:
    raw = getenv_str(name, str(default)).lower()
    return raw in {"1", "true", "yes", "y", "on"}


def split_keywords(raw: str) -> List[str]:
    pieces = re.split(r"[|；;]\s*", raw)
    return [piece.strip() for piece in pieces if piece.strip()]


def load_config(env_path: Path) -> Config:
    load_dotenv(env_path)
    default_keywords = (
        "Vibe Coding独立创业案例|Vibe Coding无代码SaaS项目|"
        "个人开发者Vibe Coding变现|轻量化AI工具Vibe Coding副业|"
        "一人AI SaaS创业Vibe Coding"
    )
    return Config(
        keywords=split_keywords(getenv_str("SEARCH_KEYWORDS", default_keywords)),
        daily_budget_cny=getenv_float("DAILY_BUDGET_CNY", 0.1),
        llm_api_key=getenv_str("LLM_API_KEY"),
        llm_base_url=getenv_str("LLM_BASE_URL", "https://api.deepseek.com"),
        llm_model=getenv_str("LLM_MODEL", "deepseek-chat"),
        llm_input_price_cny_per_1m_tokens=getenv_float(
            "LLM_INPUT_PRICE_CNY_PER_1M_TOKENS", 1.0
        ),
        llm_output_price_cny_per_1m_tokens=getenv_float(
            "LLM_OUTPUT_PRICE_CNY_PER_1M_TOKENS", 2.0
        ),
        search_provider=getenv_str("SEARCH_PROVIDER", "duckduckgo").lower(),
        serper_api_key=getenv_str("SERPER_API_KEY"),
        max_results_per_keyword=getenv_int("MAX_RESULTS_PER_KEYWORD", 6),
        max_pages_per_keyword=getenv_int("MAX_PAGES_PER_KEYWORD", 8),
        max_source_chars=getenv_int("MAX_SOURCE_CHARS", 3500),
        max_llm_output_tokens=getenv_int("MAX_LLM_OUTPUT_TOKENS", 1600),
        request_timeout_seconds=getenv_int("REQUEST_TIMEOUT_SECONDS", 20),
        output_dir=Path(getenv_str("OUTPUT_DIR", "outputs")),
        run_time=getenv_str("RUN_TIME", "09:00"),
        exchange_rates=getenv_str("EXCHANGE_RATES", "USD:7.2,EUR:7.8,GBP:9.2"),
        enable_enrichment=getenv_bool("ENABLE_ENRICHMENT", True),
        enrichment_results_per_case=getenv_int("ENRICHMENT_RESULTS_PER_CASE", 4),
        enrichment_source_chars=getenv_int("ENRICHMENT_SOURCE_CHARS", 1800),
        min_completeness_score=getenv_float("MIN_COMPLETENESS_SCORE", 0.82),
        rejected_output=getenv_bool("WRITE_REJECTED_CANDIDATES", True),
    )


class TokenBudget:
    """Persistent daily token budget, measured in CNY."""

    def __init__(self, config: Config, state_dir: Path, run_date: str) -> None:
        self.config = config
        self.state_dir = state_dir
        self.state_dir.mkdir(parents=True, exist_ok=True)
        self.state_path = state_dir / f"budget_{run_date}.json"
        self.spent_cny = 0.0
        self.estimated_tokens = 0
        self.actual_tokens = 0
        self.events: List[Dict[str, Any]] = []
        self.stopped = False
        self._load()

    def _load(self) -> None:
        if not self.state_path.exists():
            return
        try:
            data = json.loads(self.state_path.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            return
        self.spent_cny = float(data.get("spent_cny", 0.0))
        self.estimated_tokens = int(data.get("estimated_tokens", 0))
        self.actual_tokens = int(data.get("actual_tokens", 0))
        self.events = list(data.get("events", []))

    def _save(self) -> None:
        data = {
            "spent_cny": round(self.spent_cny, 8),
            "daily_budget_cny": self.config.daily_budget_cny,
            "estimated_tokens": self.estimated_tokens,
            "actual_tokens": self.actual_tokens,
            "events": self.events[-100:],
        }
        self.state_path.write_text(
            json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8"
        )

    def remaining_cny(self) -> float:
        return max(0.0, self.config.daily_budget_cny - self.spent_cny)

    def estimate_cost(self, input_tokens: int, output_tokens: int) -> float:
        input_cost = (
            input_tokens
            / 1_000_000
            * self.config.llm_input_price_cny_per_1m_tokens
        )
        output_cost = (
            output_tokens
            / 1_000_000
            * self.config.llm_output_price_cny_per_1m_tokens
        )
        return input_cost + output_cost

    def can_afford(self, input_tokens: int, output_tokens: int) -> bool:
        estimated_cost = self.estimate_cost(input_tokens, output_tokens)
        if self.spent_cny + estimated_cost > self.config.daily_budget_cny:
            self.stopped = True
            self.events.append(
                {
                    "time": dt.datetime.now().isoformat(timespec="seconds"),
                    "event": "budget_stop_before_call",
                    "estimated_cost_cny": round(estimated_cost, 8),
                    "spent_cny": round(self.spent_cny, 8),
                }
            )
            self._save()
            return False
        return True

    def record_call(
        self,
        *,
        keyword: str,
        url: str,
        estimated_input_tokens: int,
        reserved_output_tokens: int,
        usage: Optional[Dict[str, int]],
    ) -> None:
        if usage:
            input_tokens = int(usage.get("prompt_tokens", estimated_input_tokens))
            output_tokens = int(usage.get("completion_tokens", 0))
            total_tokens = int(usage.get("total_tokens", input_tokens + output_tokens))
        else:
            input_tokens = estimated_input_tokens
            output_tokens = reserved_output_tokens
            total_tokens = input_tokens + output_tokens

        cost = self.estimate_cost(input_tokens, output_tokens)
        self.spent_cny += cost
        self.estimated_tokens += estimated_input_tokens + reserved_output_tokens
        self.actual_tokens += total_tokens
        self.events.append(
            {
                "time": dt.datetime.now().isoformat(timespec="seconds"),
                "event": "llm_call",
                "keyword": keyword,
                "url": url,
                "input_tokens": input_tokens,
                "output_tokens": output_tokens,
                "cost_cny": round(cost, 8),
                "spent_cny": round(self.spent_cny, 8),
            }
        )
        if self.spent_cny >= self.config.daily_budget_cny:
            self.stopped = True
        self._save()


def approx_tokens(text: str) -> int:
    # Conservative mixed Chinese/English estimate for pre-call budget gating.
    return max(1, int(len(text) / 1.5) + 32)


def search_web(keyword: str, config: Config) -> List[Dict[str, str]]:
    if config.search_provider == "serper":
        if not config.serper_api_key:
            print("SERPER_API_KEY 为空，跳过 Serper 搜索。", file=sys.stderr)
            return []
        response = requests.post(
            "https://google.serper.dev/search",
            headers={
                "X-API-KEY": config.serper_api_key,
                "Content-Type": "application/json",
            },
            json={"q": keyword, "num": config.max_results_per_keyword, "hl": "zh-cn"},
            timeout=config.request_timeout_seconds,
        )
        response.raise_for_status()
        data = response.json()
        return [
            {
                "title": item.get("title", ""),
                "href": item.get("link", ""),
                "body": item.get("snippet", ""),
            }
            for item in data.get("organic", [])
            if item.get("link")
        ]

    results: List[Dict[str, str]] = []
    with DDGS() as ddgs:
        for item in ddgs.text(
            keyword,
            region="wt-wt",
            safesearch="moderate",
            max_results=config.max_results_per_keyword,
        ):
            href = item.get("href") or item.get("url") or ""
            if href:
                results.append(
                    {
                        "title": item.get("title", ""),
                        "href": href,
                        "body": item.get("body", ""),
                    }
                )
    return results


def fetch_page_text(url: str, config: Config) -> str:
    try:
        response = requests.get(
            url,
            headers={
                "User-Agent": (
                    "Mozilla/5.0 (compatible; VibeCaseCollector/1.0; "
                    "+https://example.local)"
                )
            },
            timeout=config.request_timeout_seconds,
        )
        response.raise_for_status()
    except requests.RequestException as exc:
        return f"页面抓取失败：{exc}"

    content_type = response.headers.get("content-type", "").lower()
    if "text" not in content_type and "html" not in content_type and "json" not in content_type:
        return f"页面不是文本内容：{content_type or MISSING}"

    soup = BeautifulSoup(response.text, "html.parser")
    for tag in soup(["script", "style", "noscript", "svg"]):
        tag.decompose()
    title = soup.title.get_text(" ", strip=True) if soup.title else ""
    text = soup.get_text("\n", strip=True)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return f"{title}\n\n{text}".strip()


def build_extraction_prompt(
    *,
    keyword: str,
    search_result: Dict[str, str],
    page_text: str,
    config: Config,
) -> str:
    source_text = page_text[: config.max_source_chars]
    return f"""
你是专业、保守的事实核验型资料抽取助手。只能依据下方“公开来源材料”抽取真实已落地案例，不得使用常识补全，不得编造。

任务范围：
- 仅收录 Vibe Coding / AI 辅助编程 / 无代码或低代码方式完成的个人或小团队 AI 工具创业真实落地案例。
- 纯概念、教程文章、工具清单、没有产品名称或没有公开来源链接的内容，返回空数组。
- 任一字段没有公开可查资料，必须填“未披露”。
- 同一字段出现矛盾数值，字段值以“存疑：...”开头，并列出全部不同数据。
- 禁止使用“推断/可能/猜测/似乎”等表述填关键字段；如果只能推断，必须填“未披露”。

字段约束：
- 编程背景只能是：有、无、自学、未披露。
- 产品类型只能是：SaaS订阅工具、本地客户端、网页插件、API服务、开源免费工具、未披露。
- 收入模式只能使用这些表达或组合：SaaS月订阅、一次性买断、API按量收费、广告变现、企业定制服务、未披露。
- 月收入必须换算为人民币并标注原始货币，例如“5万元人民币（$7200美元）”；若没有月收入公开数据，填“未披露”。汇率参考：{config.exchange_rates}。
- 数据可信度只能是：★、★★、★★★、★★★★。
- 是否单人项目只能写“是，团队规模：...”或“否，团队规模：...”或“未披露”。
- 解决痛点最多3条。
- 风险点必须包含：合规风险、平台依赖风险、市场竞争风险、长期运营风险。
- 机会点必须包含：横向扩展场景、纵向功能深化、B端企业转化、细分生态位卡位。
- 每个非“未披露”的关键结论都应来自材料中的明确证据；不要为了完整而推测。
- evidence_summary 写成 3-6 条“字段：证据摘要（来源URL）”，用于人工复核。
- missing_fields 列出仍然缺少公开证据的字段名；如果没有，返回空数组。
- verification_notes 写出冲突、低可信、需要人工复核的地方；没有则写“无”。

请只返回 JSON，不要 Markdown，不要解释。JSON 结构：
{{
  "cases": [
    {{
      "product_name": "产品完整对外名称",
      "founder_background": "创始人全名+职业/过往从业经历；无公开信息填未披露",
      "coding_background": "有/无/自学/未披露",
      "product_type": "SaaS订阅工具/本地客户端/网页插件/API服务/开源免费工具/未披露",
      "product_usage": "一句话精准概括软件核心功能，说明能解决用户什么操作需求",
      "target_users": "清晰划分使用群体",
      "pain_points": ["痛点1", "痛点2", "痛点3"],
      "income_model": "SaaS月订阅/一次性买断/API按量收费/广告变现/企业定制服务/未披露",
      "monthly_income": "人民币换算+原始货币/存疑/未披露",
      "ai_tools_used": "开发时用到的 Vibe Coding 工具、大模型 API；无公开资料填未披露",
      "development_time": "从启动开发到上线 MVP 总周期；无资料填未披露",
      "data_sources": "信息出处+公开可访问链接",
      "credibility": "★/★★/★★★/★★★★",
      "solo_project": "是，团队规模：1人 / 否，团队规模：... / 未披露",
      "risks": {{
        "合规风险": "未披露或具体风险",
        "平台依赖风险": "未披露或具体风险",
        "市场竞争风险": "未披露或具体风险",
        "长期运营风险": "未披露或具体风险"
      }},
      "opportunities": {{
        "横向扩展场景": "未披露或具体机会",
        "纵向功能深化": "未披露或具体机会",
        "B端企业转化": "未披露或具体机会",
        "细分生态位卡位": "未披露或具体机会"
      }},
      "evidence_summary": ["字段：证据摘要（来源URL）"],
      "missing_fields": ["缺少公开证据的字段名"],
      "verification_notes": "冲突、低可信或需要人工复核的信息；没有则写无"
    }}
  ]
}}

检索关键词：{keyword}
公开来源标题：{search_result.get("title", "")}
公开来源链接：{search_result.get("href", "")}
公开来源摘要：{search_result.get("body", "")}
公开来源材料：
{source_text}
""".strip()


def call_llm(prompt: str, config: Config) -> tuple[Dict[str, Any], Optional[Dict[str, int]]]:
    if not config.llm_api_key:
        raise RuntimeError("LLM_API_KEY 为空，无法进行结构化抽取。")

    url = f"{config.llm_base_url.rstrip('/')}/chat/completions"
    payload = {
        "model": config.llm_model,
        "messages": [
            {
                "role": "system",
                "content": "你只返回严格 JSON。不能确认的信息一律写未披露。",
            },
            {"role": "user", "content": prompt},
        ],
        "temperature": 0,
        "max_tokens": config.max_llm_output_tokens,
        "response_format": {"type": "json_object"},
    }
    response = requests.post(
        url,
        headers={
            "Authorization": f"Bearer {config.llm_api_key}",
            "Content-Type": "application/json",
        },
        json=payload,
        timeout=config.request_timeout_seconds,
    )
    response.raise_for_status()
    data = response.json()
    content = data["choices"][0]["message"]["content"]
    usage = data.get("usage")
    try:
        parsed = json.loads(content)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", content, flags=re.S)
        if not match:
            raise
        parsed = json.loads(match.group(0))
    return parsed, usage


def ensure_text(value: Any) -> str:
    if value is None:
        return MISSING
    if isinstance(value, str):
        cleaned = re.sub(r"\s+", " ", value).strip()
        return cleaned or MISSING
    if isinstance(value, (list, tuple)):
        cleaned_items = [ensure_text(item) for item in value if ensure_text(item) != MISSING]
        return "；".join(cleaned_items) if cleaned_items else MISSING
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False)
    return str(value).strip() or MISSING


def normalize_enum(value: Any, allowed: set[str]) -> str:
    text = ensure_text(value)
    return text if text in allowed else MISSING


def normalize_income_model(value: Any) -> str:
    text = ensure_text(value)
    if text == MISSING:
        return MISSING
    parts = [keyword for keyword in INCOME_MODE_KEYWORDS if keyword in text]
    return "；".join(parts) if parts else MISSING


def remove_inferred_value(value: Any) -> str:
    text = ensure_text(value)
    if text == MISSING:
        return MISSING
    if any(marker in text for marker in INFERENCE_MARKERS) and not text.startswith("存疑"):
        return MISSING
    return text


def normalize_pain_points(value: Any) -> List[str] | str:
    if value in (None, "", MISSING):
        return MISSING
    if isinstance(value, list):
        points = [ensure_text(item) for item in value]
    else:
        text = ensure_text(value)
        points = re.split(r"[；;]\s*|(?:①|②|③|\n)", text)
    points = [point for point in (p.strip(" -：:") for p in points) if point and point != MISSING]
    return points[:3] if points else MISSING


def normalize_nested(value: Any, keys: Iterable[str]) -> Dict[str, str]:
    result = {key: MISSING for key in keys}
    if isinstance(value, dict):
        for key in keys:
            result[key] = ensure_text(value.get(key, MISSING))
    return result


def normalize_list(value: Any) -> List[str]:
    if value in (None, "", MISSING):
        return []
    if isinstance(value, list):
        items = [ensure_text(item) for item in value]
    else:
        items = re.split(r"[；;\n]+", ensure_text(value))
    cleaned = [item.strip(" -") for item in items if item.strip(" -") and item != MISSING]
    return [
        item
        for item in cleaned
        if not any(marker in item for marker in INFERENCE_MARKERS)
    ]


def remove_fields_with_inferred_evidence(case: Dict[str, Any]) -> None:
    evidence = "；".join(case.get("evidence_summary", []))
    checks = {
        "product_type": ("产品类型", "产品形态"),
        "income_model": ("收入模式", "变现模式"),
        "monthly_income": ("月收入", "MRR", "ARR"),
        "ai_tools_used": ("AI工具", "使用的AI工具"),
        "development_time": ("开发时间", "上线周期", "MVP"),
        "solo_project": ("单人项目", "团队规模"),
    }
    for field, labels in checks.items():
        for label in labels:
            if re.search(rf"{re.escape(label)}[^；。\n]*({'|'.join(INFERENCE_MARKERS)})", evidence):
                case[field] = MISSING
                break


def has_public_link(text: str) -> bool:
    return bool(re.search(r"https?://", text))


def normalize_product_key(name: str) -> str:
    return re.sub(r"[^a-z0-9\u4e00-\u9fff]+", "", name.lower())


def is_missing_value(value: Any) -> bool:
    if value in (None, "", MISSING):
        return True
    if isinstance(value, str):
        return value.strip() in {"", MISSING}
    if isinstance(value, list):
        return len(value) == 0
    if isinstance(value, dict):
        return all(is_missing_value(item) for item in value.values())
    return False


def missing_field_names(case: Dict[str, Any]) -> List[str]:
    checks = {
        "创始人姓名+背景": case.get("founder_background"),
        "编程背景": case.get("coding_background"),
        "产品类型": case.get("product_type"),
        "产品用途": case.get("product_usage"),
        "目标人群": case.get("target_users"),
        "解决痛点": case.get("pain_points"),
        "收入模式": case.get("income_model"),
        "月收入": case.get("monthly_income"),
        "使用的AI工具": case.get("ai_tools_used"),
        "开发时间": case.get("development_time"),
        "是否单人项目": case.get("solo_project"),
    }
    return [name for name, value in checks.items() if is_missing_value(value)]


def completeness_score(case: Dict[str, Any]) -> float:
    weights = {
        "product_name": 1.0,
        "founder_background": 1.0,
        "coding_background": 0.8,
        "product_type": 0.8,
        "product_usage": 1.0,
        "target_users": 0.8,
        "pain_points": 0.8,
        "income_model": 1.0,
        "monthly_income": 1.0,
        "ai_tools_used": 1.0,
        "development_time": 1.0,
        "data_sources": 1.0,
        "credibility": 0.6,
        "solo_project": 0.8,
    }
    total = sum(weights.values())
    got = 0.0
    for field, weight in weights.items():
        value = case.get(field)
        if field == "data_sources":
            present = not is_missing_value(value) and has_public_link(ensure_text(value))
        else:
            present = not is_missing_value(value)
        if present:
            got += weight
    return round(got / total, 3)


def sanitize_case(raw: Dict[str, Any], case_id: int, fallback_url: str) -> Optional[Dict[str, Any]]:
    product_name = ensure_text(raw.get("product_name"))
    data_sources = ensure_text(raw.get("data_sources"))
    if product_name == MISSING:
        return None
    if not has_public_link(data_sources):
        if fallback_url:
            data_sources = f"{data_sources if data_sources != MISSING else '检索结果'}：{fallback_url}"
        else:
            return None

    credibility = ensure_text(raw.get("credibility"))
    if credibility not in TRUST_LEVELS:
        credibility = MISSING

    solo_project = ensure_text(raw.get("solo_project"))
    if solo_project != MISSING and not (solo_project.startswith("是") or solo_project.startswith("否")):
        solo_project = MISSING

    sanitized = {
        "case_id": case_id,
        "product_name": product_name,
        "founder_background": remove_inferred_value(raw.get("founder_background")),
        "coding_background": normalize_enum(raw.get("coding_background"), CODING_BACKGROUNDS),
        "product_type": normalize_enum(raw.get("product_type"), PRODUCT_TYPES),
        "product_usage": remove_inferred_value(raw.get("product_usage")),
        "target_users": remove_inferred_value(raw.get("target_users")),
        "pain_points": normalize_pain_points(raw.get("pain_points")),
        "income_model": normalize_income_model(raw.get("income_model")),
        "monthly_income": remove_inferred_value(raw.get("monthly_income")),
        "ai_tools_used": remove_inferred_value(raw.get("ai_tools_used")),
        "development_time": remove_inferred_value(raw.get("development_time")),
        "data_sources": data_sources,
        "credibility": credibility,
        "solo_project": solo_project,
        "risks": normalize_nested(raw.get("risks"), RISK_KEYS),
        "opportunities": normalize_nested(raw.get("opportunities"), OPPORTUNITY_KEYS),
        "evidence_summary": normalize_list(raw.get("evidence_summary")),
        "missing_fields": normalize_list(raw.get("missing_fields")),
        "completeness_score": 0.0,
        "verification_notes": ensure_text(raw.get("verification_notes", "无")),
    }
    calculated_missing = missing_field_names(sanitized)
    if not sanitized["missing_fields"]:
        sanitized["missing_fields"] = calculated_missing
    else:
        merged_missing = list(dict.fromkeys([*sanitized["missing_fields"], *calculated_missing]))
        sanitized["missing_fields"] = merged_missing
    remove_fields_with_inferred_evidence(sanitized)
    sanitized["missing_fields"] = missing_field_names(sanitized)
    sanitized["completeness_score"] = completeness_score(sanitized)
    return sanitized


def extract_cases_from_result(
    *,
    keyword: str,
    result: Dict[str, str],
    config: Config,
    budget: TokenBudget,
) -> List[Dict[str, Any]]:
    page_text = fetch_page_text(result["href"], config)
    prompt = build_extraction_prompt(
        keyword=keyword,
        search_result=result,
        page_text=page_text,
        config=config,
    )
    input_tokens = approx_tokens(prompt)
    output_tokens = config.max_llm_output_tokens
    if not budget.can_afford(input_tokens, output_tokens):
        return []

    parsed, usage = call_llm(prompt, config)
    budget.record_call(
        keyword=keyword,
        url=result["href"],
        estimated_input_tokens=input_tokens,
        reserved_output_tokens=output_tokens,
        usage=usage,
    )
    raw_cases = parsed.get("cases", [])
    if not isinstance(raw_cases, list):
        return []
    sanitized: List[Dict[str, Any]] = []
    for raw in raw_cases:
        if isinstance(raw, dict):
            item = sanitize_case(raw, case_id=0, fallback_url=result["href"])
            if item:
                sanitized.append(item)
    return sanitized


def source_block(result: Dict[str, str], page_text: str, max_chars: int) -> str:
    return "\n".join(
        [
            f"标题：{result.get('title', '')}",
            f"链接：{result.get('href', '')}",
            f"摘要：{result.get('body', '')}",
            "正文摘录：",
            page_text[:max_chars],
        ]
    ).strip()


def enrichment_queries(case: Dict[str, Any]) -> List[str]:
    product = case["product_name"]
    founder = case.get("founder_background", "")
    founder_hint = "" if founder == MISSING else founder.split("，", 1)[0]
    return [
        f'"{product}" founder revenue MRR AI coding',
        f'"{product}" "Vibe Coding" "Cursor" "Claude"',
        f'"{product}" pricing revenue launch founder',
        f'"{product}" {founder_hint} startup interview' if founder_hint else f'"{product}" startup interview',
    ]


def build_enrichment_prompt(
    *,
    existing_case: Dict[str, Any],
    sources: List[str],
    config: Config,
) -> str:
    source_text = "\n\n--- SOURCE ---\n\n".join(sources)
    return f"""
你是事实核验型创业案例研究员。请基于“已有案例草稿”和“追加公开来源材料”做二次补全与纠错。

目标：
- 提高字段完整度，但只能填入材料中有明确证据的信息。
- 如果材料没有证据，继续填“未披露”，不要猜测。
- 禁止用“推断/可能/猜测/似乎”等不确定表述填关键字段；只能推断时填“未披露”。
- 若追加来源推翻已有字段，请以追加来源为准，并在 verification_notes 说明。
- 若数值冲突，字段值以“存疑：”开头并列出全部数值和来源。
- 必须保留公开链接，data_sources 尽量列出多个来源。
- evidence_summary 必须写 3-8 条，格式为“字段：证据摘要（来源URL）”。
- missing_fields 必须列出仍缺公开证据的关键字段。

字段约束：
- 编程背景只能是：有、无、自学、未披露。
- 产品类型只能是：SaaS订阅工具、本地客户端、网页插件、API服务、开源免费工具、未披露。
- 收入模式只能使用这些表达或组合：SaaS月订阅、一次性买断、API按量收费、广告变现、企业定制服务、未披露。
- 月收入必须换算为人民币并标注原始货币；没有公开月收入填“未披露”。汇率参考：{config.exchange_rates}。
- 数据可信度只能是：★、★★、★★★、★★★★。
- 是否单人项目只能写“是，团队规模：...”或“否，团队规模：...”或“未披露”。
- 风险点和机会点可以基于公开资料及产品客观依赖做保守归纳；无依据填“未披露”。

只返回 JSON，不要 Markdown。JSON 结构：
{{
  "case": {{
    "product_name": "产品完整对外名称",
    "founder_background": "创始人全名+背景/未披露",
    "coding_background": "有/无/自学/未披露",
    "product_type": "SaaS订阅工具/本地客户端/网页插件/API服务/开源免费工具/未披露",
    "product_usage": "一句话精准概括核心功能",
    "target_users": "目标人群",
    "pain_points": ["痛点1", "痛点2", "痛点3"],
    "income_model": "收入模式/未披露",
    "monthly_income": "人民币换算+原始货币/存疑/未披露",
    "ai_tools_used": "开发时使用的AI工具/未披露",
    "development_time": "启动到上线MVP周期/未披露",
    "data_sources": "来源名称+链接，多个来源用；分隔",
    "credibility": "★/★★/★★★/★★★★",
    "solo_project": "是，团队规模：1人 / 否，团队规模：... / 未披露",
    "risks": {{
      "合规风险": "未披露或具体风险",
      "平台依赖风险": "未披露或具体风险",
      "市场竞争风险": "未披露或具体风险",
      "长期运营风险": "未披露或具体风险"
    }},
    "opportunities": {{
      "横向扩展场景": "未披露或具体机会",
      "纵向功能深化": "未披露或具体机会",
      "B端企业转化": "未披露或具体机会",
      "细分生态位卡位": "未披露或具体机会"
    }},
    "evidence_summary": ["字段：证据摘要（来源URL）"],
    "missing_fields": ["仍缺公开证据的字段名"],
    "verification_notes": "冲突、低可信或需要人工复核的信息；没有则写无"
  }}
}}

已有案例草稿：
{json.dumps(existing_case, ensure_ascii=False, indent=2)}

追加公开来源材料：
{source_text}
""".strip()


def call_llm_with_budget(
    *,
    prompt: str,
    config: Config,
    budget: TokenBudget,
    keyword: str,
    url: str,
) -> Optional[Dict[str, Any]]:
    input_tokens = approx_tokens(prompt)
    output_tokens = config.max_llm_output_tokens
    if not budget.can_afford(input_tokens, output_tokens):
        return None
    parsed, usage = call_llm(prompt, config)
    budget.record_call(
        keyword=keyword,
        url=url,
        estimated_input_tokens=input_tokens,
        reserved_output_tokens=output_tokens,
        usage=usage,
    )
    return parsed


def enrich_case(
    *,
    case: Dict[str, Any],
    config: Config,
    budget: TokenBudget,
) -> Dict[str, Any]:
    if not config.enable_enrichment or budget.stopped:
        return case

    sources: List[str] = []
    seen_urls: set[str] = set()
    for query in enrichment_queries(case):
        if budget.stopped or len(sources) >= config.enrichment_results_per_case:
            break
        try:
            results = search_web(query, config)
        except Exception as exc:  # noqa: BLE001
            print(f"补充检索失败：{query}；原因：{exc}", file=sys.stderr)
            continue
        for result in results:
            url = result.get("href", "")
            if (
                not url
                or url in seen_urls
                or urlparse(url).scheme not in {"http", "https"}
            ):
                continue
            seen_urls.add(url)
            page_text = fetch_page_text(url, config)
            sources.append(source_block(result, page_text, config.enrichment_source_chars))
            if len(sources) >= config.enrichment_results_per_case:
                break

    if not sources:
        return case

    prompt = build_enrichment_prompt(existing_case=case, sources=sources, config=config)
    try:
        parsed = call_llm_with_budget(
            prompt=prompt,
            config=config,
            budget=budget,
            keyword=f"enrich:{case['product_name']}",
            url=";".join(sorted(seen_urls)[:3]),
        )
    except Exception as exc:  # noqa: BLE001
        print(f"案例补全失败：{case['product_name']}；原因：{exc}", file=sys.stderr)
        return case

    if not parsed or not isinstance(parsed.get("case"), dict):
        return case
    enriched = sanitize_case(parsed["case"], case.get("case_id", 0), case.get("data_sources", ""))
    return enriched or case


def format_pain_points(value: List[str] | str) -> str:
    if value == MISSING or not value:
        return MISSING
    assert isinstance(value, list)
    labels = ("①", "②", "③")
    return " ".join(f"{labels[idx]}{point}" for idx, point in enumerate(value[:3]))


def format_nested(title: str, data: Dict[str, str], keys: Iterable[str]) -> str:
    lines = [f"{title}："]
    for key in keys:
        lines.append(f"   - {key}：{data.get(key, MISSING)}")
    return "\n".join(lines)


def format_list_field(value: List[str] | str) -> str:
    if not value or value == MISSING:
        return MISSING
    if isinstance(value, str):
        return value
    return "；".join(value) if value else MISSING


def format_case(case: Dict[str, Any]) -> str:
    lines = [
        f"1. 案例编号：{case['case_id']}",
        f"2. 产品名称：{case['product_name']}",
        f"3. 创始人姓名+背景：{case['founder_background']}",
        f"4. 编程背景：{case['coding_background']}",
        f"5. 产品类型：{case['product_type']}",
        f"6. 产品用途：{case['product_usage']}",
        f"7. 目标人群：{case['target_users']}",
        f"8. 解决痛点：{format_pain_points(case['pain_points'])}",
        f"9. 收入模式：{case['income_model']}",
        f"10. 月收入：{case['monthly_income']}",
        f"11. 使用的AI工具：{case['ai_tools_used']}",
        f"12. 开发时间：{case['development_time']}",
        f"13. 数据来源+链接：{case['data_sources']}",
        f"14. 数据可信度（星级标准，仅4档可选）：{case['credibility']}",
        f"15. 是否单人项目：{case['solo_project']}",
        format_nested("16. 风险点（四类各1条，缺一不可）", case["risks"], RISK_KEYS),
        format_nested("17. 机会点（四类各1条，缺一不可）", case["opportunities"], OPPORTUNITY_KEYS),
        f"18. 公开证据摘要：{format_list_field(case.get('evidence_summary', []))}",
        f"19. 仍缺失字段：{format_list_field(case.get('missing_fields', []))}",
        f"20. 数据完整度评分：{case.get('completeness_score', 0):.0%}",
        f"21. 复核说明：{case.get('verification_notes', MISSING)}",
    ]
    return "\n".join(lines)


def write_txt(
    cases: List[Dict[str, Any]],
    path: Path,
    run_date: str,
    budget: TokenBudget,
    rejected_cases: Optional[List[Dict[str, Any]]] = None,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    header = [
        f"Vibe Coding独立创业案例自动采集报告",
        f"检索日期：{run_date}",
        f"预算上限：{budget.config.daily_budget_cny}元人民币",
        f"已记录API花费：{budget.spent_cny:.6f}元人民币",
        "",
    ]
    if not cases:
        header.append("本次未检索到同时满足“真实落地、公开可查证、字段可溯源”的案例。")
    body = "\n\n".join(format_case(case) for case in cases)
    rejected_body = ""
    if rejected_cases:
        rejected_body = "\n\n--- 信息不足候选，未进入正式案例 ---\n\n" + "\n\n".join(
            format_case(case) for case in rejected_cases
        )
    path.write_text("\n".join(header) + body + rejected_body + "\n", encoding="utf-8")


def write_docx(
    cases: List[Dict[str, Any]],
    path: Path,
    run_date: str,
    budget: TokenBudget,
    rejected_cases: Optional[List[Dict[str, Any]]] = None,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading("Vibe Coding独立创业案例自动采集报告", 0)
    document.add_paragraph(f"检索日期：{run_date}")
    document.add_paragraph(f"预算上限：{budget.config.daily_budget_cny}元人民币")
    document.add_paragraph(f"已记录API花费：{budget.spent_cny:.6f}元人民币")
    if not cases:
        document.add_paragraph("本次未检索到同时满足“真实落地、公开可查证、字段可溯源”的案例。")
    for case in cases:
        document.add_heading(f"案例 {case['case_id']}：{case['product_name']}", level=1)
        for line in format_case(case).splitlines():
            document.add_paragraph(line)
    if rejected_cases:
        document.add_heading("信息不足候选，未进入正式案例", level=1)
        for case in rejected_cases:
            document.add_heading(f"候选：{case['product_name']}", level=2)
            for line in format_case(case).splitlines():
                document.add_paragraph(line)
    document.save(path)


def collect_cases(
    config: Config, run_date: str, max_cases: Optional[int]
) -> tuple[List[Dict[str, Any]], List[Dict[str, Any]], TokenBudget]:
    budget = TokenBudget(config, Path(".run_state"), run_date)
    all_cases: List[Dict[str, Any]] = []
    rejected_cases: List[Dict[str, Any]] = []
    seen: set[str] = set()
    pages_examined = 0

    if not config.llm_api_key:
        print("LLM_API_KEY 为空：将只生成空报告，不调用付费API。", file=sys.stderr)
        return all_cases, rejected_cases, budget

    for keyword in config.keywords:
        if budget.stopped:
            print("已达到或接近单日预算上限，停止后续关键词检索。", file=sys.stderr)
            break
        print(f"检索关键词：{keyword}")
        try:
            results = search_web(keyword, config)
        except Exception as exc:  # noqa: BLE001 - keep unattended runs alive
            print(f"关键词检索失败：{keyword}；原因：{exc}", file=sys.stderr)
            continue

        for result in results:
            if budget.stopped:
                print("已达到或接近单日预算上限，停止后续页面抽取。", file=sys.stderr)
                break
            if pages_examined >= config.max_pages_per_keyword * max(1, len(config.keywords)):
                break
            url = result.get("href", "")
            if not url or urlparse(url).scheme not in {"http", "https"}:
                continue
            pages_examined += 1
            try:
                extracted = extract_cases_from_result(
                    keyword=keyword,
                    result=result,
                    config=config,
                    budget=budget,
                )
            except Exception as exc:  # noqa: BLE001 - log and continue
                print(f"页面抽取失败：{url}；原因：{exc}", file=sys.stderr)
                continue

            for item in extracted:
                key = normalize_product_key(item["product_name"])
                if not key or key in seen:
                    continue
                seen.add(key)
                item = enrich_case(case=item, config=config, budget=budget)
                item["completeness_score"] = completeness_score(item)
                item["missing_fields"] = missing_field_names(item)
                if item["completeness_score"] < config.min_completeness_score:
                    item["case_id"] = len(rejected_cases) + 1
                    rejected_cases.append(item)
                    print(
                        f"跳过信息不足候选：{item['product_name']} "
                        f"（完整度 {item['completeness_score']:.0%}）"
                    )
                    continue
                item["case_id"] = len(all_cases) + 1
                all_cases.append(item)
                print(
                    f"收录案例：{item['case_id']} {item['product_name']} "
                    f"（完整度 {item['completeness_score']:.0%}）"
                )
                if max_cases and len(all_cases) >= max_cases:
                    return all_cases, rejected_cases, budget

    return all_cases, rejected_cases, budget


def parse_run_time(value: str) -> dt.time:
    try:
        hour, minute = value.split(":", 1)
        return dt.time(hour=int(hour), minute=int(minute))
    except Exception as exc:  # noqa: BLE001
        raise ValueError("RUN_TIME 必须是 HH:MM 格式，例如 09:00") from exc


def seconds_until_next_run(run_time: dt.time) -> float:
    now = dt.datetime.now()
    target = dt.datetime.combine(now.date(), run_time)
    if target <= now:
        target += dt.timedelta(days=1)
    return max(1.0, (target - now).total_seconds())


def run_once(args: argparse.Namespace) -> None:
    env_path = Path(args.env)
    config = load_config(env_path)
    run_date = args.date or dt.date.today().isoformat()
    cases, rejected_cases, budget = collect_cases(config, run_date, args.max_cases)

    output_dir = config.output_dir
    txt_path = output_dir / f"vibe_coding_cases_{run_date}.txt"
    docx_path = output_dir / f"vibe_coding_cases_{run_date}.docx"
    rejected_for_output = rejected_cases if config.rejected_output else None
    write_txt(cases, txt_path, run_date, budget, rejected_for_output)
    write_docx(cases, docx_path, run_date, budget, rejected_for_output)
    print(f"TXT已生成：{txt_path}")
    print(f"DOCX已生成：{docx_path}")


def run_loop(args: argparse.Namespace) -> None:
    config = load_config(Path(args.env))
    run_time = parse_run_time(config.run_time)
    print(f"常驻定时模式已启动，每天 {config.run_time} 自动运行。按 Ctrl+C 停止。")
    while True:
        wait_seconds = seconds_until_next_run(run_time)
        print(f"距离下次运行约 {int(wait_seconds)} 秒。")
        time.sleep(wait_seconds)
        run_once(args)


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="采集公开可查证的 Vibe Coding 独立创业案例，并输出 TXT + DOCX。"
    )
    parser.add_argument("--env", default=".env", help="配置文件路径，默认 .env")
    parser.add_argument("--date", default="", help="输出日期，默认今天，格式 YYYY-MM-DD")
    parser.add_argument("--max-cases", type=int, default=0, help="最多收录案例数，0 表示不限")
    parser.add_argument("--loop", action="store_true", help="常驻定时模式，每天按 .env 的 RUN_TIME 运行")
    return parser


def main() -> int:
    parser = build_arg_parser()
    args = parser.parse_args()
    if args.max_cases <= 0:
        args.max_cases = None
    if args.loop:
        run_loop(args)
    else:
        run_once(args)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
